import datetime
import json
from pathlib import Path

import openpyxl
import pytest

from engram.wikiskill import spreadsheet as s


def _wb(path: Path, sheets: dict[str, dict[str, object]]) -> Path:
    """sheets: {name: {"A1": value, ...}} → an .xlsx at path (first dict entry = first sheet)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for ref, v in cells.items():
            ws[ref] = v
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --- grader: port of the official compare_workbooks ------------------------------------------

@pytest.mark.parametrize("gold,got,ok", [
    (5, 5, True),
    (1.234, 1.2349, True),              # rounded to 2 decimals
    (1.23, 1.24, False),
    (5, "5", True),                     # numeric strings are compared as numbers (official transform)
    ("abc", "abd", False),
    (5, "x", False),                    # type mismatch after transform
    ("", None, True),                   # empty string ≡ empty cell
    (None, None, True),
    (datetime.datetime(2024, 1, 1), 45292, True),   # datetime → Excel serial, rounded to the day
])
def test_grade_cell_values(tmp_path, gold, got, ok):
    g = _wb(tmp_path / "g.xlsx", {"S": {"A1": gold}})
    o = _wb(tmp_path / "o.xlsx", {"S": {"A1": got}})
    score, msg = s.grade(g, o, "A1")
    assert score == (1.0 if ok else 0.0)
    assert ok or "A1" in msg


def test_grade_ranges_and_sheet_names(tmp_path):
    g = _wb(tmp_path / "g.xlsx", {"First": {"A1": 1}, "My Sheet": {"A1": 1, "A2": 2, "C1": "x"}})
    o = _wb(tmp_path / "o.xlsx", {"First": {"A1": 1}, "My Sheet": {"A1": 1, "A2": 2, "C1": "x"}})
    assert s.grade(g, o, "'My Sheet'!A1:A2,'My Sheet'!C1") == (1.0, "")
    o2 = _wb(tmp_path / "o2.xlsx", {"First": {"A1": 1}, "My Sheet": {"A1": 1, "A2": 3, "C1": "x"}})
    score, msg = s.grade(g, o2, "'My Sheet'!A1:A2,'My Sheet'!C1")
    assert score == 0.0 and "A2" in msg and "My Sheet" in msg


def test_grade_position_without_sheet_uses_first_golden_sheet(tmp_path):
    g = _wb(tmp_path / "g.xlsx", {"Main": {"B2": 7}, "Other": {"B2": 0}})
    o = _wb(tmp_path / "o.xlsx", {"Main": {"B2": 7}, "Other": {"B2": 99}})
    assert s.grade(g, o, "B2") == (1.0, "")


def test_grade_missing_sheet_or_file(tmp_path):
    g = _wb(tmp_path / "g.xlsx", {"Main": {"A1": 1}})
    o = _wb(tmp_path / "o.xlsx", {"Renamed": {"A1": 1}})
    assert s.grade(g, o, "Main!A1") == (0.0, "worksheet not found: Main")
    assert s.grade(g, tmp_path / "missing.xlsx", "A1") == (0.0, "File not exist")
    (tmp_path / "broken.xlsx").write_text("not a zip")
    score, msg = s.grade(g, tmp_path / "broken.xlsx", "A1")
    assert score == 0.0 and msg.startswith("unreadable output")


def test_grade_formula_without_cached_value_reads_as_empty(tmp_path):
    """No recalculation (official protocol): openpyxl writes formulas without cached values."""
    g = _wb(tmp_path / "g.xlsx", {"S": {"A1": 2, "A2": 3, "A3": 5}})
    o = _wb(tmp_path / "o.xlsx", {"S": {"A1": 2, "A2": 3, "A3": "=A1+A2"}})
    score, msg = s.grade(g, o, "A3")
    assert score == 0.0 and "None" in msg


# --- preview / prepare / prompt / sandbox ----------------------------------------------------

def test_preview_first_five_rows_of_every_sheet(tmp_path):
    rows = {f"A{i}": f"r{i}" for i in range(1, 9)}
    p = _wb(tmp_path / "x.xlsx", {"One": rows, "Two": {"B1": 3.5}})
    text = s.preview(p)
    assert "Sheet: One" in text and "Sheet: Two" in text
    assert "r5" in text and "r6" not in text and "3.5" in text


def test_preview_is_capped(tmp_path):
    p = _wb(tmp_path / "x.xlsx", {"Wide": {f"{openpyxl.utils.get_column_letter(c)}1": "v" * 40 for c in range(1, 200)}})
    assert len(s.preview(p)) <= s.PREVIEW_CAP + 100


def _ws(tmp_path):
    ws = tmp_path / "ws"
    data = ws / ".data" / s.DATA_DIRNAME / "spreadsheet" / "t1"
    _wb(data / "1_t1_init.xlsx", {"S": {"A1": "in"}})
    _wb(data / "1_t1_golden.xlsx", {"S": {"A1": "gold"}})
    (ws / ".venv/bin").mkdir(parents=True)
    return ws


TASK = {"id": "t1", "instruction": "Put gold in A1", "instruction_type": "Cell-Level Manipulation",
        "answer_position": "A1", "answer": "A1"}


def test_prepare_stages_only_the_input(tmp_path):
    ws = _ws(tmp_path)
    wd = ws / "work/iter-1/t1"
    wd.mkdir(parents=True)
    s.SpreadsheetBench().prepare(ws, TASK, wd)
    assert sorted(p.name for p in wd.iterdir()) == ["input.xlsx"]  # never the golden file


def test_user_prompt_fills_the_e1_fields(tmp_path):
    ws = _ws(tmp_path)
    wd = ws / "work/iter-1/t1"
    wd.mkdir(parents=True)
    b = s.SpreadsheetBench()
    b.prepare(ws, TASK, wd)
    up = b.user_prompt(TASK, wd)
    for field in (f"working_directory: {wd}", "instruction: Put gold in A1", f"spreadsheet_path: {wd / 'input.xlsx'}",
                  "instruction_type: Cell-Level Manipulation", "answer_position: A1",
                  f"output_path: {wd / 'output.xlsx'}", "spreadsheet_content:", "Sheet: S"):
        assert field in up
    sp = b.system_prompt("## Skills\nX")
    assert "spreadsheet expert" in sp and "## Skills\nX" in sp and "{skill_section}" not in sp


def test_claude_opts_sandbox_env_turns(tmp_path, monkeypatch):
    ws = _ws(tmp_path)
    wd = ws / "work/iter-1/t1"
    monkeypatch.setattr(s, "_python_root", lambda venv: Path("/uv/python/cpython-3.13"))
    o = s.SpreadsheetBench().claude_opts(ws, wd)
    assert o["tools"] == ["Bash"] and o["allowed_tools"] == ["Bash"] and o["stream"] is True and o["max_turns"] == 30
    sb = o["settings"]["sandbox"]
    assert sb["enabled"] is True and sb["allowUnsandboxedCommands"] is False
    assert sb["filesystem"]["denyRead"] == ["~/"]
    assert sb["filesystem"]["allowRead"] == [str(wd), str(ws / ".venv"), "/uv/python/cpython-3.13"]
    assert o["env"]["PATH"].startswith(f"{ws / '.venv/bin'}:") and o["env"]["VIRTUAL_ENV"] == str(ws / ".venv")


def test_score_grades_output_against_golden(tmp_path):
    ws = _ws(tmp_path)
    wd = ws / "work/iter-1/t1"
    b = s.SpreadsheetBench()
    wd.mkdir(parents=True)
    assert b.score(ws, TASK, "done", wd) == (0.0, "File not exist")
    _wb(wd / "output.xlsx", {"S": {"A1": "gold"}})
    assert b.score(ws, TASK, "done", wd) == (1.0, "")


def test_init_splits_and_builds_the_agent_venv(tmp_path, monkeypatch):
    ws = tmp_path / "ws"
    src = tmp_path / "extracted" / s.DATA_DIRNAME
    (src / "spreadsheet").mkdir(parents=True)
    (src / "dataset.json").write_text(json.dumps([
        {"id": f"t{i}", "instruction": "i", "instruction_type": "Cell-Level Manipulation", "answer_position": "A1",
         "spreadsheet_path": f"spreadsheet/t{i}"} for i in range(400)]))
    monkeypatch.setattr(s.SpreadsheetBench, "_download", lambda self, ws_: src)
    made = {}
    monkeypatch.setattr(s, "_make_venv", lambda venv: made.setdefault("venv", venv))
    splits = s.SpreadsheetBench().init(ws, seed=0)
    assert {k: len(v) for k, v in splits.items()} == {"train": 80, "val": 40, "test": 280}
    t = splits["train"][0]
    assert set(t) == {"id", "instruction", "instruction_type", "answer_position", "answer"}
    assert t["answer"] == t["answer_position"]
    assert made["venv"] == ws / ".venv"


def test_grade_malformed_positions_follow_the_official_grader(tmp_path):
    """Verified-400 has 4 odd positions; the official parser + its bare `except` decide their verdict.
    Whole-column ranges ("A:G") make the official parser raise → always 0."""
    g = _wb(tmp_path / "g.xlsx", {"S": {"A1": 1}})
    score, msg = s.grade(g, g, "'S'!A:G")
    assert score == 0.0 and msg.startswith("grader error")


def test_init_normalizes_integer_ids_to_strings(tmp_path, monkeypatch):
    """Verified-400 task 45944 has an int id; paths and trace file names need str."""
    ws = tmp_path / "ws"
    src = tmp_path / "extracted" / s.DATA_DIRNAME
    src.mkdir(parents=True)
    (src / "dataset.json").write_text(json.dumps([
        {"id": i if i % 2 else f"t{i}", "instruction": "i", "instruction_type": "Cell-Level Manipulation",
         "answer_position": "A1", "spreadsheet_path": "x"} for i in range(400)]))
    monkeypatch.setattr(s.SpreadsheetBench, "_download", lambda self, ws_: src)
    monkeypatch.setattr(s, "_make_venv", lambda venv: None)
    splits = s.SpreadsheetBench().init(ws, seed=0)
    assert all(isinstance(t["id"], str) for v in splits.values() for t in v)


@pytest.mark.parametrize("init_name,golden_name", [
    ("1_t1_init.xlsx", "1_t1_golden.xlsx"),       # the documented layout
    ("initial.xlsx", "golden.xlsx"),              # 5 Verified-400 tasks (e.g. 13284)
    ("1_t1_init.xlsx", "1_t9_golden.xlsx"),       # 42930 ships its golden as 1_43930_golden.xlsx
])
def test_task_files_resolve_the_packaging_variants(tmp_path, init_name, golden_name):
    ws = tmp_path / "ws"
    d = ws / ".data" / s.DATA_DIRNAME / "spreadsheet" / "t1"
    _wb(d / init_name, {"S": {"A1": "in"}})
    _wb(d / golden_name, {"S": {"A1": "gold"}})
    wd = ws / "work/iter-1/t1"
    wd.mkdir(parents=True)
    b = s.SpreadsheetBench()
    b.prepare(ws, TASK, wd)
    assert openpyxl.load_workbook(wd / "input.xlsx")["S"]["A1"].value == "in"
    _wb(wd / "output.xlsx", {"S": {"A1": "gold"}})
    assert b.score(ws, TASK, "", wd) == (1.0, "")


def test_missing_input_fails_loud(tmp_path):
    ws = tmp_path / "ws"
    (ws / ".data" / s.DATA_DIRNAME / "spreadsheet" / "t1").mkdir(parents=True)
    with pytest.raises(FileNotFoundError, match="t1"):
        s.SpreadsheetBench().prepare(ws, TASK, tmp_path)


def test_python_root_covers_the_whole_symlink_chain(tmp_path):
    """Live smoke: uv venvs point at `…/python/cpython-3.13-…/bin/python3.13`, itself a symlink to
    `cpython-3.13.13-…`. Re-opening only the resolved install dir left the intermediate link under the
    denied ~/, so the sandbox made the venv's python3 unusable and `python3` fell through to
    /usr/bin/python3 (no openpyxl). The allowed root must be the directory holding the installs."""
    root = tmp_path / "uv" / "python"
    real = root / "cpython-3.13.13-macos"
    (real / "bin").mkdir(parents=True)
    (real / "bin" / "python3.13").write_text("")
    (root / "cpython-3.13-macos").symlink_to(real)
    venv = tmp_path / "ws" / ".venv"
    (venv / "bin").mkdir(parents=True)
    (venv / "bin" / "python").symlink_to(root / "cpython-3.13-macos" / "bin" / "python3.13")
    assert s._python_root(venv) == root
