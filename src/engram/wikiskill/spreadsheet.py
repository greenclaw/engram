"""SpreadsheetBench (Verified 400) as a WikiSkill bench: the agent edits a real .xlsx with Python in a
sandboxed bash shell (E.1 SpreadsheetBench prompt, Table 6: 80 / 40 / 280), graded by a port of the
official `compare_workbooks` — cached cell values only, no formula recalculation."""
from __future__ import annotations

import datetime
import json
import os
import shutil
import subprocess
import tarfile
from pathlib import Path

import openpyxl

from engram.wikiskill.bench import PROMPTS, make_splits

HF_REPO, TARBALL = "KAKA22/SpreadsheetBench", "spreadsheetbench_verified_400.tar.gz"
DATA_DIRNAME = "spreadsheetbench_verified_400"
PREVIEW_ROWS, PREVIEW_CELL, PREVIEW_CAP = 5, 60, 6000
MAX_TURNS = 30  # SkillOpt's SpreadsheetBench setting: "up to 30 turns"
SESSION_TIMEOUT = 1800  # s; a timeout is a retried failure, so it must exceed any honest 30-turn session


# --- grader (official evaluation.py semantics) -----------------------------------------------

def _transform(v):
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    if isinstance(v, datetime.time):
        return str(v)[:-3]
    if isinstance(v, datetime.datetime):
        delta = v - datetime.datetime(1899, 12, 30)
        return round(delta.days + delta.seconds / 86400.0, 0)
    if isinstance(v, str):
        try:
            return round(float(v), 2)
        except ValueError:
            return v
    return v


def _same(v1, v2) -> bool:
    v1, v2 = _transform(v1), _transform(v2)
    if v1 in ("", None) and v2 in ("", None):
        return True
    return type(v1) is type(v2) and v1 == v2


# Range parsing ported verbatim from the official grader (not openpyxl's range_boundaries): four
# Verified-400 positions are malformed ("A:G", "BD2:308", spaces after commas, commas inside a sheet
# name) and the official verdict on them is whatever this parser + its bare `except` produce.
def _col_name2num(name: str) -> int:
    num = 0
    for c in name:
        num = num * 26 + (ord(c) - ord("A") + 1)
    return num


def _col_num2name(n: int) -> str:
    name = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        name = chr(65 + rem) + name
    return name


def _split_ref(ref: str) -> tuple[int, int]:
    col = "".join(ch for ch in ref if not ch.isdigit())
    row = "".join(ch for ch in ref if ch.isdigit())
    return _col_name2num(col), int(row)  # int("") raises on whole-column refs, as in the original


def _cells(cell_range: str) -> list[str]:
    if ":" not in cell_range:
        return [cell_range]
    start, end = cell_range.split(":")
    (c1, r1), (c2, r2) = _split_ref(start), _split_ref(end)
    return [f"{_col_num2name(c)}{r}" for c in range(c1, c2 + 1) for r in range(r1, r2 + 1)]


def grade(golden: Path, output: Path, answer_position: str) -> tuple[float, str]:
    """1.0 iff every cell of every range in `answer_position` matches the golden workbook; otherwise
    0.0 and the first difference (the feedback the trace records). A range without a sheet name
    refers to the golden file's first sheet, as in the official grader."""
    if not output.exists():
        return 0.0, "File not exist"
    try:
        wb_gt = openpyxl.load_workbook(golden, data_only=True)
        wb_out = openpyxl.load_workbook(output, data_only=True)
    except Exception as e:  # noqa: BLE001 — any unreadable output is a failed answer, not a crash
        return 0.0, f"unreadable output: {e}"[:300]
    try:
        for part in answer_position.split(","):
            sheet, rng = part.split("!") if "!" in part else (wb_gt.sheetnames[0], part)
            sheet, rng = sheet.strip("'"), rng.strip("'")
            if sheet not in wb_out:
                return 0.0, f"worksheet not found: {sheet}"
            ws_gt, ws_out = wb_gt[sheet], wb_out[sheet]
            for ref in _cells(rng):
                if not _same(ws_gt[ref].value, ws_out[ref].value):
                    return 0.0, f"{sheet}!{ref}: expected {ws_gt[ref].value!r}, got {ws_out[ref].value!r}"
    except Exception as e:  # noqa: BLE001 — the official evaluation() scores any grader exception as 0
        return 0.0, f"grader error: {type(e).__name__}: {e}"[:300]
    return 1.0, ""


# --- preview (spreadsheet_content: SpreadsheetBench's "first 5 rows" setting) -----------------

def preview(path: Path) -> str:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            out.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(max_row=PREVIEW_ROWS, values_only=True):
                out.append("\t".join("" if v is None else str(v)[:PREVIEW_CELL] for v in row))
    finally:
        wb.close()
    text = "\n".join(out)
    return text if len(text) <= PREVIEW_CAP else text[:PREVIEW_CAP] + "\n…[preview truncated]"


# --- environment ------------------------------------------------------------------------------

def _python_root(venv: Path) -> Path:
    """The directory holding the interpreter installs the venv points at (uv keeps them under ~, so
    the sandbox must re-open them for reads). Not just the resolved install dir: uv links
    `cpython-3.13-…` → `cpython-3.13.13-…`, and every link on the chain must stay readable, or the
    venv's python3 is unusable and `python3` silently falls through to /usr/bin/python3."""
    return Path(os.path.realpath(venv / "bin" / "python")).parents[2]


def _make_venv(venv: Path) -> None:
    subprocess.run(["uv", "venv", "-q", str(venv)], check=True)
    subprocess.run(["uv", "pip", "install", "-q", "--python", str(venv / "bin" / "python"),
                    "openpyxl", "pandas"], check=True)


class SpreadsheetBench:
    name = "spreadsheetbench"
    task_desc = "spreadsheet manipulation tasks by writing and running Python code (openpyxl / pandas) in a bash shell"
    SPLIT_SIZES = {"train": 80, "val": 40, "test": 280}  # Table 6

    @staticmethod
    def data_dir(ws: Path) -> Path:
        return ws / ".data" / DATA_DIRNAME

    def _download(self, ws: Path) -> Path:
        from huggingface_hub import hf_hub_download

        tar = hf_hub_download(HF_REPO, TARBALL, repo_type="dataset", cache_dir=ws / ".hf-cache")
        with tarfile.open(tar) as t:
            t.extractall(ws / ".data", filter="data")
        return self.data_dir(ws)

    def init(self, ws: Path, seed: int) -> dict[str, list[dict]]:
        src = self._download(ws)
        tasks = [{"id": str(d["id"]), "instruction": d["instruction"], "instruction_type": d["instruction_type"],
                  "answer_position": d["answer_position"], "answer": d["answer_position"]}
                 for d in json.loads((src / "dataset.json").read_text())]
        splits = make_splits(tasks, self.SPLIT_SIZES, seed)
        _make_venv(ws / ".venv")
        return splits

    def system_prompt(self, skill_section: str) -> str:
        return (PROMPTS / "spreadsheetbench.md").read_text().replace("{skill_section}", skill_section)

    def _file(self, ws: Path, task: dict, kind: str) -> Path:
        """`1_<id>_{init,golden}.xlsx`, or the packaging variants Verified-400 also ships: 5 tasks use
        `initial.xlsx` / `golden.xlsx`, and 42930's golden carries a typo'd id (`1_43930_golden.xlsx`)."""
        d = self.data_dir(ws) / "spreadsheet" / task["id"]
        candidates = [d / f"1_{task['id']}_{kind}.xlsx", d / ("initial.xlsx" if kind == "init" else "golden.xlsx"),
                      *sorted(d.glob(f"*_{kind}.xlsx"))]
        for c in candidates:
            if c.exists():
                return c
        raise FileNotFoundError(f"no {kind} spreadsheet for task {task['id']} in {d}")

    def prepare(self, ws: Path, task: dict, workdir: Path) -> None:
        shutil.copyfile(self._file(ws, task, "init"), workdir / "input.xlsx")  # never the golden file

    def user_prompt(self, task: dict, workdir: Path) -> str:
        src = workdir / "input.xlsx"
        return (f"working_directory: {workdir}\n"
                f"instruction: {task['instruction']}\n"
                f"spreadsheet_path: {src}\n"
                f"instruction_type: {task['instruction_type']}\n"
                f"answer_position: {task['answer_position']}\n"
                f"output_path: {workdir / 'output.xlsx'}\n"
                f"spreadsheet_content:\n{preview(src)}\n")

    def claude_opts(self, ws: Path, workdir: Path) -> dict:
        venv = ws / ".venv"
        sandbox = {"enabled": True, "autoAllowBashIfSandboxed": True, "allowUnsandboxedCommands": False,
                   # the paper's restriction, enforced: nothing under ~ is readable except the task
                   # workdir and the agent's Python; writes stay in cwd; network is off by default
                   "filesystem": {"denyRead": ["~/"],
                                  "allowRead": [str(workdir), str(venv), str(_python_root(venv))]}}
        return {"tools": ["Bash"], "allowed_tools": ["Bash"], "stream": True, "max_turns": MAX_TURNS,
                "timeout": SESSION_TIMEOUT,
                "settings": {"sandbox": sandbox},
                "env": {"PATH": f"{venv / 'bin'}:{os.environ.get('PATH', '')}", "VIRTUAL_ENV": str(venv)}}

    def score(self, ws: Path, task: dict, response: str, workdir: Path) -> tuple[float, str]:
        return grade(self._file(ws, task, "golden"), workdir / "output.xlsx", task["answer_position"])
